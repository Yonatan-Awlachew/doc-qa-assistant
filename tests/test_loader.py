"""Tests for reading PDF, DOCX and TXT files."""
import pytest
from docx import Document

from app.loader import load_file

SAMPLE_PDF = "samples/example_regulations.pdf"


def test_pdf_pages_and_text():
    pages = load_file(SAMPLE_PDF, "rules.pdf")
    assert len(pages) == 3
    assert pages[0]["source"] == "rules.pdf"
    assert pages[1]["page"] == 2
    assert "180 credits" in pages[1]["text"]


def test_txt_is_split_into_pages(tmp_path):
    path = tmp_path / "long.txt"
    path.write_text("word " * 2000)            # 10,000 characters
    pages = load_file(str(path), "long.txt")
    assert len(pages) == 4                     # 3000 + 3000 + 3000 + rest
    assert pages[3]["page"] == 4


def test_docx(tmp_path):
    path = tmp_path / "notes.docx"
    document = Document()
    document.add_paragraph("The exam is on Monday.")
    document.add_paragraph("Bring your student card.")
    document.save(str(path))
    pages = load_file(str(path), "notes.docx")
    assert "student card" in pages[0]["text"]


def test_unsupported_type_raises_error(tmp_path):
    path = tmp_path / "image.png"
    path.write_bytes(b"not really an image")
    with pytest.raises(ValueError):
        load_file(str(path), "image.png")